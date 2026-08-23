# -*- coding: utf-8 -*-
"""Build the Huangyan CAM Automation Studio product proposal."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "cam-automation-studio-huangyan-product-proposal.docx"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2933"
MUTED = "5D6B78"
LIGHT = "F4F6F9"
WHITE = "FFFFFF"
GRID = "CCD5DF"
EVIDENCE_FILL = "EAF4F1"
EVIDENCE_TEXT = "1F6E5C"
CAUTION_FILL = "FFF4E5"
CAUTION_TEXT = "9A5B13"
BLOCK_FILL = "FBECEC"
BLOCK_TEXT = "A33A3A"

BODY_LATIN = "Calibri"
BODY_EAST_ASIA = "Microsoft YaHei"
MONO = "Consolas"


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
    latin: str = BODY_LATIN,
    east_asia: str = BODY_EAST_ASIA,
) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, *, size: float, color: str = INK, bold: bool = False) -> None:
    style.font.name = BODY_LATIN
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_EAST_ASIA)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in CELL_MARGIN_DXA.items():
        tag = "w:start" if edge == "start" else "w:end" if edge == "end" else f"w:{edge}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(element, fill: str) -> None:
    pr = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element.get_or_add_pPr()
    shd = pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pr.append(shd)
    shd.set(qn("w:fill"), fill)


def paragraph_border(paragraph, color: str, *, size: str = "6", space: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "bottom", "left", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), space)
        node.set(qn("w:color"), color)
        p_bdr.append(node)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths: list[int], *, indent: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != PAGE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {PAGE_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        set_cant_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def format_table_paragraph(paragraph, *, header: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.12
    for run in paragraph.runs:
        set_run_font(run, size=9.25, bold=header, color=INK)


def fill_cell(cell, text: str, *, header: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_run_font(run, size=9.25, bold=header, color=INK)
    format_table_paragraph(paragraph, header=header)
    if fill:
        shade(cell._tc, fill)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        fill_cell(table.rows[0].cells[index], header, header=True, fill=LIGHT)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            fill_cell(row.cells[index], value)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_numbering_definition(doc: Document, *, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract_num.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "279")
    p_pr.append(ind)
    level.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_LATIN)
    fonts.set(qn("w:hAnsi"), BODY_LATIN)
    fonts.set(qn("w:eastAsia"), BODY_EAST_ASIA)
    r_pr.append(fonts)
    level.append(r_pr)
    abstract_num.append(level)
    numbering.append(abstract_num)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def add_bullet(doc: Document, text: str, *, bold_prefix: str | None = None) -> object:
    paragraph = doc.add_paragraph()
    apply_numbering(paragraph, doc._bullet_num_id)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, bold=True)
        run = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_numbered(doc: Document, text: str, *, compact: bool = False) -> object:
    paragraph = doc.add_paragraph()
    apply_numbering(paragraph, doc._decimal_num_id)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3 if compact else 4)
    paragraph.paragraph_format.line_spacing = 1.18 if compact else 1.208
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_body(
    doc: Document,
    text: str,
    *,
    bold_prefix: str | None = None,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    after: float = 8,
) -> object:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.333
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(rest, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, color=INK)
    return paragraph


def add_callout(
    doc: Document,
    label: str,
    text: str,
    *,
    fill: str = EVIDENCE_FILL,
    color: str = EVIDENCE_TEXT,
) -> object:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.22
    shade(paragraph._p, fill)
    paragraph_border(paragraph, color, size="4", space="6")
    run = paragraph.add_run(f"{label}  ")
    set_run_font(run, size=10.2, bold=True, color=color)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.2, color=INK)
    return paragraph


def add_code_block(doc: Document, text: str) -> object:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.14)
    paragraph.paragraph_format.right_indent = Inches(0.14)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.0
    shade(paragraph._p, LIGHT)
    paragraph_border(paragraph, GRID, size="4", space="6")
    lines = text.splitlines()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, size=8.5, color=DARK_BLUE, latin=MONO, east_asia=BODY_EAST_ASIA)
        if index != len(lines) - 1:
            run.add_break()
    return paragraph


def add_heading(
    doc: Document,
    text: str,
    *,
    level: int = 1,
    page_break_before: bool = False,
) -> object:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    return paragraph


def add_label_value(doc: Document, label: str, value: str) -> object:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(f"{label}：")
    set_run_font(run, size=10.2, bold=True, color=DARK_BLUE)
    run = paragraph.add_run(value)
    set_run_font(run, size=10.2, color=INK)
    return paragraph


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_LATIN)
    fonts.set(qn("w:hAnsi"), BODY_LATIN)
    fonts.set(qn("w:eastAsia"), BODY_EAST_ASIA)
    r_pr.extend([fonts, color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(
    doc: Document,
    code: str,
    title: str,
    publisher: str,
    published: str,
    url: str,
    note: str,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(f"[{code}] {title}。")
    set_run_font(run, size=9.2, bold=True, color=INK)
    run = paragraph.add_run(f"{publisher}，{published}。{note}  ")
    set_run_font(run, size=9.2, color=INK)
    add_hyperlink(paragraph, "打开原始来源", url)
    return paragraph


def set_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    set_style_font(styles["Normal"], size=11, color=INK)
    normal = styles["Normal"]
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    set_style_font(styles["Title"], size=26, color=INK, bold=True)
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(4)

    set_style_font(styles["Subtitle"], size=13.5, color=MUTED)
    styles["Subtitle"].paragraph_format.space_before = Pt(0)
    styles["Subtitle"].paragraph_format.space_after = Pt(8)

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        set_style_font(styles[name], size=size, color=color, bold=True)
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True

    doc._bullet_num_id = add_numbering_definition(doc, kind="bullet")
    doc._decimal_num_id = add_numbering_definition(doc, kind="decimal")

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run("CAM Automation Studio  |  黄岩模具企业产品方案")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run("内部产品提案  ·  2026-08-23  ·  第 ")
    set_run_font(run, size=8.5, color=MUTED)
    set_page_number(fp)
    run = fp.add_run(" 页")
    set_run_font(run, size=8.5, color=MUTED)

    doc.core_properties.title = "CAM Automation Studio 黄岩模具企业产品与最小版本实施方案"
    doc.core_properties.subject = "PowerMill / Siemens NX 工作流学习、审阅和自动化产品方案"
    doc.core_properties.author = "CAM Automation Studio"
    doc.core_properties.keywords = (
        "PowerMill, Siemens NX, UG, CAM, 黄岩模具, 工作流学习, 插件, Codex"
    )
    doc.core_properties.comments = "基于公开资料和当前仓库能力形成，需经工厂现场验证。"


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(22)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run("产品与试点实施方案")
    set_run_font(run, size=12, bold=True, color=MUTED)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CAM Automation Studio")
    set_run_font(run, size=26, bold=True, color=INK)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    run = subtitle.add_run("面向台州黄岩模具企业的本地优先 PowerMill / NX 工作流智能化平台")
    set_run_font(run, size=13.5, color=MUTED)

    audience = doc.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience.paragraph_format.space_after = Pt(24)
    run = audience.add_run("从专家行为记录到可审阅自动化配方，优先使用接口、命令、日志、Journal 与 API")
    set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)

    metadata = doc.add_table(rows=3, cols=4)
    pairs = [
        ("方案对象", "黄岩模具企业 / CAM 负责人", "当前底座", "仓库版本 0.5.0"),
        ("研究方法", "公开资料 + 官方接口 + 竞品", "证据日期", "截至 2026-08-23"),
        ("建议周期", "首轮 12 周", "输出性质", "MVP 与现场试点方案"),
    ]
    for row_index, row_data in enumerate(pairs):
        for col_index, value in enumerate(row_data):
            is_label = col_index in (0, 2)
            fill_cell(
                metadata.rows[row_index].cells[col_index],
                value,
                header=is_label,
                fill=LIGHT if is_label else None,
            )
    set_repeat_table_header(metadata.rows[0])
    set_table_geometry(metadata, [1200, 3480, 1200, 3480])

    add_callout(
        doc,
        "核心建议",
        "不要先做“看屏幕点击”的通用代理。先把 PowerMill 宏/命令、NX Journal/NXOpen、"
        "项目实体与执行响应接成稳定的数据闭环，再把高频专家流程升级为可安装插件。"
        "图像识别只保留为无法获得结构化数据时的辅助证据。",
    )

    add_body(
        doc,
        "本方案在当前软件已有插件化、本地采集、日志学习、多实例检测、dry-run 预览和 Codex "
        "审阅桥接能力上继续推进。第一版的价值不是直接生成机床可用 NC，而是让工厂能看见、筛选、"
        "复盘并复用高级编程人员的重复操作，同时保持目标窗口、项目版本、运行结果和仿真边界清晰。",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        after=0,
    )
    doc.add_page_break()


def build_content(doc: Document) -> None:
    add_heading(doc, "执行建议", level=1)
    add_callout(
        doc,
        "决策摘要",
        "以“本地工作台、结构化数据优先、插件按需安装、结果可预览”为产品主线。"
        "首轮 12 周只承诺读、记、筛、学、审、dry-run 六件事；目标版本 live 适配在现场副本项目中单独验收。",
    )
    add_body(
        doc,
        "黄岩已经形成大规模模具产业集群，并在 2024 年推进模具行业中小企业“N+X”数字化改造。"
        "公开资料显示，区内约有 10 万模具从业者、4000 多家模具及相关企业、180 余家设计/造型/"
        "编程服务企业，模具行业年产值接近 300 亿元；规上模具企业数字化改造覆盖率达到 85%。"
        "这些数字说明市场不是缺少单点软件，而是缺少能把分散人员、软件版本、工艺经验与审阅责任"
        "连接起来的轻量工作层。[S3]",
    )
    add_body(
        doc,
        "当前仓库 0.5.0 已具备离线学习闭环，但没有真实 PowerMill/NX 生产传输。最合理的下一步不是"
        "扩大 AI 叙事，而是完成只读原生适配、自动结构化采集和现场基线测量。所有效率提升数字都必须"
        "先作为假设，直到在 5-8 家黄岩工厂的真实编程任务中得到验证。",
    )

    add_heading(doc, "五个产品决定", level=2)
    decisions = [
        "结构化数据优先：API/project entity > 命令与响应 > PowerMill 宏/日志 > NX Journal/NXOpen 事件 > 文件与进程元数据 > OCR/图像。",
        "基础软件从零开始：默认只提供插件中心、健康状态和静态外壳；PowerMill、NX、采集、执行、Codex 审阅均按需安装。",
        "采集插件安装后自动连接并开始记录，顶部持续显示状态；高级设置可分别关闭日志记录、窗口发现和命令记录。",
        "一个主窗口解决主要任务：顶部显示 Codex、NX、PowerMill 连接摘要；点击后再展开各 PID/窗口/前台状态，避免为状态信息启动多个窗口。",
        "执行与学习分离：学习输出是配方和证据；命令运行明确选择目标实例、目标版本和测试项目，并返回仿真、碰撞/过切检查结果。",
    ]
    for item in decisions:
        add_numbered(doc, item)

    add_heading(doc, "当前底座与必须补齐的缺口", level=2)
    add_table(
        doc,
        ["能力域", "0.5.0 已具备", "首版缺口", "产品决定"],
        [
            ["插件", "空白启动、依赖、安装/卸载", "宿主侧版本适配包", "核心保持小，适配器独立发布"],
            ["采集", "本地 SQLite、增量读取、多窗口检测", "原生事件/响应与可靠会话归属", "结构化适配先行，不猜窗口归属"],
            ["学习", "ActivityEvent v1、重复序列、变量参数", "专家标签和工厂工艺词典", "人工标注专家示范，不自动推断"],
            ["执行", "命令记录、dry-run、目标实例", "PowerMill COM/.NET 与 NXOpen live 传输", "版本化、逐实例、测试项目优先"],
            ["审阅", "Codex 结构化上下文与回传", "现场审阅模板和差异对比", "AI 提建议，工程师确认结果"],
        ],
        [1300, 2550, 2550, 2960],
    )

    add_heading(doc, "成功标准", level=2)
    add_bullet(doc, "用户在 10 分钟内完成目标插件安装，并看到正确的多实例连接与自动记录状态。")
    add_bullet(doc, "手动操作、自动化操作、系统事件和命令活动可独立筛选，并能按原始事件、动作、会话、配方、运行候选五层查看。")
    add_bullet(doc, "三次以上相似专家会话可形成带来源、参数、支持度、差异和预览的配方草稿。")
    add_bullet(doc, "所有命令请求都有目标实例、请求文本、响应、耗时和结果记录，dry-run 可直接比较运行前后差异。")
    add_bullet(doc, "第一版不生成或下发机床可用 NC，重点验证工作流学习、配方复用和 CAM 项目内自动化。")

    add_heading(doc, "研究方法与证据边界", level=1, page_break_before=True)
    add_callout(
        doc,
        "重要边界",
        "本方案完成的是公开/案头研究，不声称已经访谈黄岩工厂。已验证事实来自黄岩区政府、Autodesk、"
        "Siemens 与厂商公开页面；痛点优先级、定价、效率目标和实施成本属于待现场验证的产品假设。",
        fill=CAUTION_FILL,
        color=CAUTION_TEXT,
    )
    add_body(
        doc,
        "研究使用三类证据：第一类是黄岩区政府公开的产业规模和中小企业数字化改造文件；第二类是"
        "PowerMill 与 NX 的官方自动化资料；第三类是 CAM Assist、OptiNC 等公开产品能力。搜索摘要仅"
        "用于发现线索，结论只引用已直接打开的原文。当前仓库 README、架构、测试和插件清单用于界定"
        "软件已经实现的能力。",
    )

    add_heading(doc, "黄岩产业信号", level=2)
    add_table(
        doc,
        ["已验证事实", "对产品的含义", "证据"],
        [
            ["约 10 万从业者、4000 多家相关企业、年产值近 300 亿元", "足以支撑面向区域工艺复用和服务商生态的垂直产品", "S3"],
            ["180 余家设计、造型、编程服务企业", "首批用户既可来自工厂，也可来自区域编程服务商", "S3"],
            ["规上模具企业数字化改造覆盖率 85%", "产品需兼容已有数字系统，不能要求整体替换", "S3"],
            ["2024 年计划约 10 家试点、2 家以上总承包商，形成“N+X”清单", "共性底座 + 行业/工厂插件与区域政策方向一致", "S1"],
            ["11 家试点项目完成验收，5 家成为样本企业", "试点、专家评审、现场考察和样板复制是可接受的落地路径", "S2"],
        ],
        [3200, 4800, 1360],
    )

    add_heading(doc, "由证据推导、仍需验证的需求假设", level=2)
    hypotheses = [
        "编程过程不透明：管理者知道交期和结果，但不知道时间消耗在哪些重复步骤、回滚和等待上。",
        "专家经验难迁移：高级程序员的参数选择、动作顺序和异常处理分散在个人 Journal、宏和临时习惯中。",
        "软件并行与窗口并行常见：同一操作者可能同时打开多个 NX/PowerMill 项目，自动连接必须显示目标实例，不能依赖前台窗口猜测。",
        "数字化系统已经存在但不连贯：CAM、文件、排产、质量和机床仿真之间仍需要轻量审阅和证据层。",
        "多种项目数据分散：零件、刀具、夹具、刀路和工艺参数需要在一个任务视图中串联。",
        "真正价值来自可复用流程，不是更炫的聊天：用户需要一键载入、清晰状态、参数说明和少量可预测操作。",
    ]
    for item in hypotheses:
        add_bullet(doc, item)

    add_heading(doc, "竞品给出的边界信号", level=2)
    add_table(
        doc,
        ["产品", "公开能力", "值得学习", "本项目差异"],
        [
            ["CloudNC CAM Assist", "嵌入现有 CAM，生成策略/刀路，支持 NX；厂商宣称可完成最高 80% 的 CAM 程序", "宿主内审阅、工具/材料/机床配置、先给可编辑结果", "本项目先学习工厂自己的专家行为，本地优先；效率数字不照搬"],
            ["OptiNC for PowerMill", "PowerMill 插件，100+ 功能，覆盖粗精加工、钻孔、3+2、批量参数、NC/工艺单等", "用高频垂直向导降低学习门槛，功能模块化", "先从行为挖掘确定黄岩高频模块；首版不直接输出 NC"],
        ],
        [1650, 3050, 2450, 2210],
    )
    add_body(
        doc,
        "竞品页面中的“80%”“70%”属于厂商公开宣传，不是本方案的承诺，也没有被独立验证。[S9][S10] "
        "它们只证明一个稳定的产品规律：CAM 自动化必须嵌入现有工作流、让工程师审阅，并把复杂能力"
        "包装成明确任务，而不是要求用户理解底层 AI。",
        after=0,
    )

    add_heading(doc, "目标用户与高频任务", level=1, page_break_before=True)
    add_table(
        doc,
        ["角色", "核心任务", "最怕发生", "首版价值"],
        [
            ["老板/厂长", "交期、产能、人才复制", "投入看不到收益、软件闲置", "可量化基线、按插件部署、按场景购买"],
            ["CAM 主管", "工艺标准、审阅、任务分配、异常复盘", "不同程序员结果不一致", "配方支持度、版本、参数和差异证据"],
            ["高级程序员", "复杂件策略、异常处理、带教", "记录噪声太多、整理负担过重", "自动整理会话，由本人标注专家示范"],
            ["初级程序员", "按规范完成重复建模与刀路任务", "参数选错、找不到模板、不会回滚", "步骤化配方、参数解释、dry-run 预览"],
            ["机床/质量人员", "仿真、碰撞、过切、交接", "项目结果难复现、问题定位慢", "运行差异、仿真结果和任务时间线"],
            ["IT/数字化负责人", "部署、升级、接口、性能", "版本碎片、后台失控、难排障", "插件清单、连接状态、诊断与版本矩阵"],
        ],
        [1320, 2400, 2440, 3200],
    )

    add_heading(doc, "应优先学习的工作流", level=2)
    add_body(
        doc,
        "首轮现场观察不以“功能多”为目标，而以重复频率、结果可验证性、跨程序员一致性和可测量收益排序。"
        "优先候选如下：",
    )
    priority_workflows = [
        "项目/模板初始化：单位、坐标系、毛坯、命名、目录和审阅规则。",
        "模型与参考数据导入：来源识别、尺寸检查、版本记录和稳定对象映射。",
        "刀具、刀柄、夹具和工件坐标准备：只读校验优先，异常明确提示。",
        "粗加工、半精、精加工和钻孔的重复参数设置：先学习顺序和变量，不直接假设参数正确。",
        "工作平面、边界、孔特征、刀路复制与批量参数修改：高频、易审阅、适合插件化。",
        "仿真前检查和程序交接：汇总缺少对象、未计算刀路、未审阅变更和目标版本不一致等检查项。",
    ]
    for item in priority_workflows:
        add_numbered(doc, item)

    add_heading(doc, "日志模式与层级筛选", level=2)
    add_body(
        doc,
        "用户提出的“只显示手动操作”和“按层级筛选”应成为基础交互，不是分析后的附加功能。"
        "建议把来源模式与学习层级拆成两个正交筛选器，并通过向 ActivityEvent 添加可选字段保持向后兼容。",
    )
    add_table(
        doc,
        ["筛选维度", "值", "说明"],
        [
            ["来源模式 source_mode", "manual", "用户在 PowerMill/NX 中直接操作或录制的行为"],
            ["来源模式 source_mode", "automation", "宏、Journal、插件或配方产生的动作"],
            ["来源模式 source_mode", "system", "进程、窗口、文件发现、连接与状态事件"],
            ["来源模式 source_mode", "execution_audit", "命令请求、响应、耗时和结果事件"],
            ["学习层级 view_level", "L0 原始证据", "保留原文、时间、来源行与实例"],
            ["学习层级 view_level", "L1 规范动作", "映射为 cam.*、nx.*、powermill.* 动作"],
            ["学习层级 view_level", "L2 会话", "按任务/窗口/项目组织的连续操作"],
            ["学习层级 view_level", "L3 重复配方", "跨会话出现的结构、变量和支持度"],
            ["学习层级 view_level", "L4 运行候选", "已补参数、目标版本、目标实例和预览结果的候选"],
        ],
        [2250, 2200, 4910],
    )
    add_callout(
        doc,
        "兼容性规则",
        "只新增 source_mode、expertise_label、instance_id、command_response、review_status 等可选字段；"
        "不重命名或删除 ActivityEvent v1 现有字段。PowerMill 解析进入独立适配器，不向 NX Journal 解析器混入语法。",
    )

    add_heading(doc, "产品定义与交互原则", level=1, page_break_before=True)
    add_callout(
        doc,
        "一句话定位",
        "面向黄岩模具企业的本地 CAM 工作流副驾驶：从 PowerMill/NX 的结构化行为中学习专家经验，"
        "生成可回放、可参数化、可审阅、可定向 dry-run 的自动化配方。",
    )
    add_heading(doc, "用户第一次打开软件", level=2)
    first_run = [
        "只看到插件中心、软件健康状态和“未安装任何业务模块”的清晰空状态。",
        "安装 PowerMill 或 UG/NX 工作流插件后，工作台出现对应导入、连接和分析入口。",
        "安装本地行为记录插件后立即自动连接并开始后台记录；顶部常驻“记录中”状态，不再弹出多步设置窗口。",
        "高级设置中可分别关闭日志记录、窗口发现和命令记录，也可一键暂停全部记录。",
        "安装命令中心后可选择目标实例并运行 dry-run；目标版本适配包决定可用的查询和命令类型。",
    ]
    for item in first_run:
        add_numbered(doc, item)

    add_heading(doc, "一个窗口内的主要导航", level=2)
    add_table(
        doc,
        ["区域", "默认信息密度", "主要操作"],
        [
            ["顶部连接条", "Codex、NX、PowerMill 三项摘要；实例数和异常色", "点击展开实例抽屉"],
            ["插件中心", "名称、用途、依赖、版本、状态", "安装、卸载、更新、查看说明"],
            ["工作台", "输入、日志、会话、配方、审阅五个页签", "导入、筛选、标注、学习、导出"],
            ["高级设置", "记录开关、路径、性能、版本、诊断", "分类开关、重扫、导出、缓存维护"],
            ["连接抽屉", "产品、PID、窗口标题、前台、日志映射", "选定 dry-run 目标并查看连接响应"],
        ],
        [1700, 3700, 3960],
    )
    add_body(
        doc,
        "连接状态必须在首页可见，而不是藏在独立诊断页。用户同时打开多个软件窗口时，摘要显示"
        "“NX 2”“PowerMill 3”；展开后逐项列出 instance_id、PID、窗口标题、项目和前台标记。"
        "日志源无法可靠映射到某一窗口时显示“未映射”，由用户在同一抽屉内校正，不额外打开状态窗口。",
    )

    add_heading(doc, "插件目录", level=2)
    add_table(
        doc,
        ["模块", "首版状态", "职责", "主要输入/输出"],
        [
            ["基础内核", "内置", "健康、插件目录、静态 UI、安装注册表", "插件清单、状态、任务"],
            ["UG / NX 工作流", "按需安装", "Journal 静态解析、重复流程、预览", "Journal / JSONL、配方"],
            ["PowerMill 工作流", "按需安装", "宏/命令日志学习、参数化配方", ".mac / 日志、配方"],
            ["本地行为记录", "按需安装", "后台增量读取、多实例、筛选、导出", "日志、进程、窗口、本地库"],
            ["命令中心", "按需安装", "目标实例、dry-run、响应、任务历史", "CAM 命令、DiffReport"],
            ["Codex 审阅", "按需安装", "结构化上下文、发现项回传", "上下文包、修改建议"],
            ["版本化原生适配包", "后续安装", "PowerMill COM/.NET 或 NXOpen 目标版本传输", "版本 API、响应"],
        ],
        [2100, 1550, 3600, 2110],
    )

    add_heading(doc, "结构化数据优先级", level=2)
    levels = [
        ("P1", "原生 API 与项目实体", "对象、刀具、刀路、工作平面、选择集、状态和稳定 ID"),
        ("P2", "命令与响应", "执行文本、返回值、错误、耗时、目标实例和版本"),
        ("P3", "PowerMill 宏/日志", ".mac、RecordMacro 和命令轨迹；适合快速覆盖真实手工步骤"),
        ("P4", "NX Journal / NXOpen 事件", "记录调用、对象构建器、参数和撤销边界；需去除录制噪声"),
        ("P5", "文件与进程元数据", "路径、时间、哈希、PID、窗口状态和来源映射"),
        ("P6", "图像/OCR 辅助", "仅在没有接口时帮助定位 UI 状态，不作为项目对象和命令语义的主来源"),
    ]
    add_table(doc, ["优先级", "数据面", "用途"], [list(item) for item in levels], [900, 2700, 5760])

    add_heading(doc, "数据契约与系统架构", level=1, page_break_before=True)
    add_body(
        doc,
        "PowerMill 与 NX 使用不同自动化语言，但学习问题相同：把一个会话中的有序动作、参数、来源、"
        "时间、实例、参数和结果映射为统一事件。共享动作使用 cam.*；产品专有动作分别使用 "
        "powermill.* 或 nx.*。适配器互不解析对方语法。",
    )
    add_code_block(
        doc,
        """PowerMill .mac / API / response ----> PowerMill adapter ----┐
                                                              │
NX Journal / NXOpen event ----------> NX adapter -------------+--> ActivityEvent v1
                                                              │          |
Process / window / file metadata ----> capture adapter -------┘          v
                                                               SQLite WAL + event index
                                                                          |
                                                           sessionization + sequence mining
                                                                          |
                                                       reviewed recipe + evidence + dry-run
                                                                          |
                                         target-version adapter + simulation + result report""",
    )

    add_heading(doc, "建议的 ActivityEvent 可选扩展", level=2)
    add_table(
        doc,
        ["字段", "类型", "用途", "兼容策略"],
        [
            ["source_mode", "enum?", "manual / automation / system / execution_audit", "可选；缺失时 unknown"],
            ["expertise_label", "enum?", "routine / expert_demo / training / unknown", "仅人工标注"],
            ["instance_id", "string?", "产品 + PID + 窗口句柄的实例标识", "用于连接和会话归属"],
            ["command_response", "object?", "状态、文本摘要、耗时、错误码", "只在有原生传输时写入"],
            ["review_status", "enum?", "draft / reviewed / needs_changes / ready_for_test", "显示配方整理进度"],
            ["target_version", "string?", "PowerMill/NX 精确版本", "运行候选必填"],
            ["recipe_hash", "string?", "锁定已审阅配方版本", "用于差异比较和结果复现"],
        ],
        [2050, 1350, 3700, 2260],
    )

    add_heading(doc, "学习流水线", level=2)
    pipeline = [
        "采集：增量 tail 受支持日志，或接收 Journal/API 事件；以内容哈希和游标去重。",
        "规范化：保留引号内字面量，映射为 cam.* / nx.* / powermill.* 动作并保留来源行。",
        "会话化：优先使用实例、项目和显式标记；证据不足时保持未映射，不强行合并。",
        "挖掘：按结构签名发现重复序列，把跨会话变化的字面量提升为类型化参数。",
        "审阅：显示支持度、来源、参数、目标版本和差异；专家确认动作含义。",
        "预览：生成 PowerMill 宏草稿或 NXOpen 非修改预览，不自动发往宿主。",
        "运行候选：完成 Codex/人工审阅后绑定目标实例和测试项目，生成 dry-run、仿真和结果报告。",
    ]
    for item in pipeline:
        add_numbered(doc, item, compact=True)

    add_heading(doc, "PowerMill 原生适配设计", level=1, page_break_before=True)
    add_body(
        doc,
        "Autodesk 官方资料确认 PowerMill 可以录制用户操作为宏；录制只保存用户在对话框中改变的值，"
        "所以采集侧必须保留“未改变的默认值可能没有出现”这一证据限制。[S4] Autodesk 官方 GitHub "
        "还提供 PowerMill 自动化库和插件示例，可作为外部 COM/.NET 宿主与进程内插件的实现参考。[S5][S6]",
    )
    add_table(
        doc,
        ["阶段", "传输", "产品能力", "版本验证重点"],
        [
            ["PM-0 离线", ".mac / .log / JSONL", "解析、筛选、学习、生成审阅草稿", "语法、参数、来源行"],
            ["PM-1 只读", "PMAutomation / DoCommandEx 查询", "连接、项目/实体查询、响应、RecordMacro", "版本、实例选择、线程与超时"],
            ["PM-2 目标 dry-run", "PMAutomation 或进程内 QueueCommand", "已审阅命令、明确目标实例", "测试项目、快照、回滚、响应"],
            ["PM-3 工艺辅助", "版本化适配器", "钻孔、粗加工、精加工等场景插件", "仿真、碰撞、项目结果；仍不自动 NC"],
        ],
        [1450, 2500, 2770, 2640],
    )
    add_heading(doc, "PowerMill 适配器合同", level=2)
    add_code_block(
        doc,
        """discover_instances() -> [InstanceDescriptor]
connect(target_version, target_instance_id, read_only=True)
query(command) -> CommandResponse
record_macro(path) -> RecordingHandle
get_project_snapshot() -> SnapshotMetadata
preview(reviewed_recipe, parameters) -> DiffReport
queue(reviewed_command, run_context) -> CommandResponse
disconnect()""",
    )
    add_bullet(doc, "每个实例独立串行命令队列；UI 线程只接收状态，不等待长计算。")
    add_bullet(doc, "DoCommand/DoCommandEx/QueueCommand 的能力差异必须按目标版本实机测试，官方示例维护优先级较低。")
    add_bullet(doc, "read_only=True 时仅暴露项目和实体查询；查询方法由版本适配包维护。")
    add_bullet(doc, "所有 live 请求需要明确 target_instance_id；界面不会把前台窗口自动当作目标。")

    add_heading(doc, "Siemens NX / UG 原生适配设计", level=1, page_break_before=True)
    add_body(
        doc,
        "Siemens 官方 NX Open Python 参考和程序员指南按 NX 系列发布，当前已核验页面分别标识 NX 2406 "
        "与 NX 2212 系列。[S7][S8] 这直接支持“按目标版本适配和验证”的工程策略：不能把一个版本录制的 "
        "Journal 当作所有版本都稳定的生产接口。",
    )
    add_table(
        doc,
        ["阶段", "输入/传输", "允许能力", "关键处理"],
        [
            ["NX-0 离线", "NX Open Python Journal / JSONL", "AST 静态解析、动作映射、配方和预览", "绝不 import 或执行 Journal"],
            ["NX-1 只读", "目标版本 NXOpen Python", "会话、部件、对象和 CAM 设置查询", "对照本机 UGOPEN/pythonStubs"],
            ["NX-2 目标 dry-run", "NXOpen Builder / 任务脚本", "在测试项目预览对象和参数变化", "稳定选择器、Undo Mark、快照、响应"],
            ["NX-3 工艺辅助", "脚本/插件包", "钻孔、工序、刀具等场景向导", "Teamcenter/许可/批处理/仿真验证"],
        ],
        [1450, 2500, 2770, 2640],
    )
    add_heading(doc, "NX 稳定性规则", level=2)
    nx_rules = [
        "先录制一个最小 Journal，确认单一操作的真实 NXOpen 调用，再去除噪声。",
        "FindObject 录制标识只作为证据；生产实现优先改为受控名称、属性、PMI 或几何查询。",
        "规划与修改分开：preview 报告将选择的对象和参数，运行任务显示目标实例、版本和项目。",
        "以目标 NX 安装自带 Python stubs 为最终合同；社区代码只用于理解形态，不能替代版本验证。",
        "NX 专属解析只保留在 nx_journal.py；PowerMill 语法进入独立 powermill_macro.py 或现有 PowerMill 适配层。",
    ]
    for item in nx_rules:
        add_bullet(doc, item)

    add_heading(doc, "性能与可靠性预算", level=2)
    add_callout(
        doc,
        "指标性质",
        "以下是首版工程验收目标，不是已测结果。现有仓库只证明 20,000 条命令解析在保守 3 秒测试上限内；"
        "真实宿主开销必须在黄岩工厂目标版本和项目规模上测量。",
        fill=CAUTION_FILL,
        color=CAUTION_TEXT,
    )
    add_table(
        doc,
        ["路径", "首版目标", "降级策略"],
        [
            ["后台日志发现", "新增事件 2 秒内可见；平均 CPU < 2%", "退避扫描、只 tail 增量、可暂停"],
            ["本地存储", "100,000 事件导入 < 10 秒；UI 查询 p95 < 200 ms", "批量事务、WAL、索引、分页"],
            ["宿主查询", "单命令超时可配置，默认 5 秒", "断路、重连、实例标红、不阻塞其他实例"],
            ["工作流学习", "10,000 会话内可重复、结果确定", "限制候选长度、后台任务、保留进度"],
            ["多实例", "每实例独立队列和健康状态", "单实例故障隔离，不切换目标"],
            ["恢复", "游标/哈希去重，重启不重复写入", "校验点、任务事件、手动重扫"],
        ],
        [2000, 3900, 3460],
    )

    add_heading(doc, "12 周最小版本", level=1, page_break_before=True)
    add_callout(
        doc,
        "MVP 定义",
        "交付一个能在 PowerMill 与 NX 并行环境中自动发现连接、后台记录结构化行为、按模式/层级筛选、"
        "学习重复专家流程、生成可审阅配方并对指定实例执行 dry-run 的本地插件平台。",
    )
    add_table(
        doc,
        ["周", "交付", "现场验证"],
        [
            ["1-2", "5-8 家工厂访谈与编程观察；版本/日志来源清单；基线 KPI", "至少 15 个真实会话样本"],
            ["3-4", "PowerMill 只读连接 POC；宏录制与响应证据", "2 个目标版本、双窗口连接不串线"],
            ["3-5", "NX 只读连接 POC；Journal 与目标 stubs 对照", "2 个目标版本、稳定对象查询样例"],
            ["5-7", "采集服务增强：source_mode、层级、实例、自动连接、记录开关", "手动/自动/命令筛选准确率人工抽检"],
            ["7-9", "会话化、专家标签、重复配方、差异和参数视图", "3 类高频流程形成可复盘候选"],
            ["9-10", "Codex 审阅包、目标实例 dry-run、响应和任务历史", "超时、断连和异常恢复演练"],
            ["11-12", "两家样板工厂试点、培训、性能/收益测量、发布包", "Go/No-Go 评审和下一版范围"],
        ],
        [900, 4850, 3610],
    )

    add_heading(doc, "首版包含", level=2)
    included = [
        "一个主窗口、插件中心、顶部多软件连接摘要和实例抽屉。",
        "PowerMill 宏/日志/JSONL 与 NX Journal/JSONL 离线学习。",
        "本地记录插件安装后默认启动；提供分类开关、暂停、导出、重扫和缓存统计。",
        "manual / automation / system / execution_audit 模式筛选与 L0-L4 层级筛选。",
        "人工标注“日常操作/专家示范/培训”，不根据姓名或动作自动判定专家。",
        "重复序列、参数候选、支持度、来源、版本和差异报告。",
        "PowerMill 宏草稿、NX 非修改预览、Codex 审阅上下文和目标实例 dry-run。",
    ]
    for item in included:
        add_bullet(doc, item)

    add_heading(doc, "首版聚焦边界", level=2)
    excluded = [
        "主要识别源是 API、命令、响应、宏、Journal 和项目实体；OCR 只补充无法由接口表达的界面状态。",
        "会话归属依据实例、项目和日志映射；无法确认时保留为“未映射”，交给用户在连接抽屉内校正。",
        "首版目标是从记录中发现重复步骤并形成配方，不承诺自动生成完整加工策略或最优循环时间。",
        "首版运行输出是项目内 dry-run、差异和仿真结果，不生成、后处理或发送机床可用 NC/G-code。",
        "live 能力按 PowerMill/NX 版本适配包逐步增加，每个版本都要用真实项目样本验证。",
    ]
    for item in excluded:
        add_bullet(doc, item)

    add_heading(doc, "最小演示脚本", level=2)
    demo_steps = [
        "启动基础软件，确认没有业务模块；从插件中心安装 UG/NX、PowerMill 和本地行为记录。",
        "安装后自动显示 NX 2 个窗口、PowerMill 2 个窗口和 Codex 状态；记录插件自动开始，展开实例详情。",
        "导入或自动采集三次相似手工会话，筛选“仅手动 + L2 会话”，将其中两次标为专家示范。",
        "生成重复配方，查看变量参数、来源行、支持度、差异与 PowerMill/NX 预览。",
        "导出 Codex 审阅上下文，回传 findings、参数建议和待确认项。",
        "选择一个明确 target_instance_id，在测试项目上提交 dry-run；展示响应、耗时和前后差异。",
        "在高级设置关闭窗口发现并暂停记录，再恢复记录并确认增量游标可以续接。",
    ]
    for item in demo_steps:
        add_numbered(doc, item)

    add_heading(doc, "产品功能模块详解", level=1, page_break_before=True)
    add_callout(
        doc,
        "产品主流程",
        "连接软件 → 自动记录 → 筛选手工会话 → 发现重复流程 → 编辑参数配方 → "
        "Codex 审阅 → 指定 NX/PowerMill 实例 dry-run → 对比结果。"
        "所有步骤都在一个主窗口中完成，详情用抽屉或页签展开。",
    )

    add_heading(doc, "功能优先级矩阵", level=2)
    add_table(
        doc,
        ["功能模块", "用户任务", "MVP 交付", "验收信号"],
        [
            ["多软件连接中心", "同时看清 NX、PowerMill、Codex 与多个窗口", "自动发现、实例抽屉、项目/版本/响应", "10 秒内显示全部实例"],
            ["日志中心", "只看手工或指定层级的记录", "模式筛选、L0-L4、搜索、保存视图", "筛选结果与样本一致"],
            ["会话时间线", "复盘一次编程任务及回滚", "动作序列、耗时、项目切换、会话对比", "可解释会话差异"],
            ["工作流学习", "发现专家反复执行的步骤", "重复序列、支持度、变量、候选排序", "产出 3 类候选流程"],
            ["配方编辑器", "把候选流程整理成可复用模板", "步骤、参数、条件、版本、预览", "可从来源追到每一步"],
            ["PowerMill 适配器", "读取宏、命令和项目实体", "宏/日志解析、只读查询、响应映射", "双实例不串线"],
            ["NX 适配器", "读取 Journal 和 NXOpen 对象", "AST 解析、只读查询、稳定对象提示", "目标版本样本通过"],
            ["命令中心", "向指定实例提交查询或 dry-run", "目标选择、队列、响应、DiffReport", "任务结果可复现"],
            ["Codex 审阅助手", "快速解释动作、参数和差异", "上下文包、发现项、修改建议回传", "配方修改有依据"],
            ["插件中心", "按场景安装软件能力", "安装、依赖、更新、版本适配包", "空白内核可独立启动"],
            ["诊断与统计", "定位断连、慢任务和采集缺口", "性能图、任务耗时、游标、重连、导出", "故障可在一页定位"],
        ],
        [1800, 2650, 3160, 1750],
    )

    add_heading(doc, "核心页面流程", level=2)
    add_code_block(
        doc,
        """首次启动
  -> 插件中心：安装“PowerMill 学习”或“UG/NX 学习”
  -> 自动连接：顶部显示 PowerMill 2 / NX 3 / Codex 1
  -> 自动记录：日志页实时增加事件
  -> 筛选：manual + L2 会话 + 当前项目
  -> 对比：选择 3 个相似会话
  -> 学习：生成重复流程候选和变量参数
  -> 配方：调整步骤、条件、参数默认值
  -> 审阅：Codex 返回解释、差异和待确认项
  -> 命令中心：选择目标实例与测试项目
  -> dry-run：返回响应、耗时、对象变化和仿真摘要""",
    )
    add_body(
        doc,
        "首页默认只呈现连接摘要、正在记录的事件数、最近会话和待整理配方。用户点击一个摘要后在右侧"
        "抽屉查看实例详情；点击一条会话后进入同一工作台的时间线页签，不创建新的顶级窗口。",
    )

    add_heading(doc, "模块交互图", level=2)
    add_code_block(
        doc,
        """PowerMill instances                         NX instances
  macro / command / API response                  Journal / NXOpen
              |                                         |
              v                                         v
     PowerMill adapter                            NX adapter
              |                                         |
              +------------> ActivityEvent v1 <----------+
                           |        |
                           |        +--> connection center
                           v
                 local event store + index
                           |
             +-------------+--------------+
             |             |              |
          log center   session timeline   diagnostics
             |             |
             +------> workflow learning
                           |
                    recipe editor
                      |          |
                      v          v
                Codex review   command center
                                   |
                        target instance + dry-run
                                   |
                         response + DiffReport""",
    )

    add_heading(doc, "连接、日志与会话", level=2, page_break_before=True)
    add_table(
        doc,
        ["模块", "默认界面", "关键交互", "后台实现"],
        [
            ["多软件连接中心", "顶部三项摘要 + 右侧实例抽屉", "按产品/版本筛选，固定目标，手动校正未映射会话", "进程发现、版本探测、心跳、每实例状态机"],
            ["日志中心", "筛选条 + 虚拟滚动表格 + 详情面板", "模式、L0-L4、时间、项目、实例、动作、全文搜索", "SQLite WAL、增量游标、索引、分页"],
            ["会话时间线", "动作轨道 + 耗时分布 + 详情", "合并/拆分会话，标记专家示范，对比 2-5 次会话", "会话化、时间间隔、项目和实例特征"],
        ],
        [1900, 2850, 3070, 1540],
    )
    add_bullet(doc, "记录插件安装后默认自动连接并开始记录；状态条显示“记录中 / 已暂停 / 源中断 / 正在追赶”。")
    add_bullet(doc, "模式筛选可单选“仅手动”，也可多选 manual + automation；层级筛选独立工作，可保存为常用视图。")
    add_bullet(doc, "当同时打开多个 PowerMill/NX 窗口时，每条事件都显示实例色标；未映射事件单独成组。")

    add_heading(doc, "学习、配方与 Codex", level=2)
    add_table(
        doc,
        ["模块", "输入", "主要结果", "用户可调整"],
        [
            ["工作流学习", "2 个以上会话或一个选择集", "重复步骤、支持度、变量参数、流程分支", "最小支持度、忽略动作、会话范围"],
            ["配方编辑器", "重复流程候选", "步骤列表、参数表、适用条件、版本范围、预览", "顺序、启停、参数类型、默认值、说明"],
            ["Codex 审阅助手", "配方 + 来源事件 + 对比结果", "动作解释、差异摘要、参数建议、待确认项", "采用建议、补充上下文、重新生成"],
        ],
        [2050, 2700, 3000, 1610],
    )
    add_body(
        doc,
        "配方编辑器是首版最重要的产物页面。左侧按步骤显示 cam.* / nx.* / powermill.* 动作，中间编辑"
        "参数和条件，右侧始终显示来源证据与版本差异。Codex 只接收结构化审阅包，返回 findings、"
        "suggested_parameters 和 questions，用户在同一页面逐项采用。",
    )

    add_heading(doc, "适配器与命令中心", level=2)
    add_table(
        doc,
        ["模块", "首版能力", "后续增强", "性能设计"],
        [
            ["PowerMill 适配器", ".mac/.log 导入、RecordMacro 发现、DoCommandEx 查询与响应", "版本化 COM/.NET、QueueCommand、项目实体快照", "每实例串行队列，查询超时独立"],
            ["NX 适配器", "Journal AST、JSONL、目标 stubs 对照、对象查询", "NXOpen Builder、Teamcenter 场景、稳定对象选择器", "Journal 离线解析，宿主查询异步"],
            ["命令中心", "查询命令、配方预览、目标实例 dry-run、DiffReport", "按场景增加项目内动作和批处理任务", "后台任务队列、取消、进度和结果缓存"],
        ],
        [2050, 3250, 2800, 1260],
    )
    add_body(
        doc,
        "命令中心明确显示产品、PID、项目、版本、配方哈希和参数。用户可直接输入受适配器支持的查询"
        "命令，也可从配方创建任务。运行后在同一页面看到原始响应、结构化结果、耗时、对象变化和仿真摘要。"
        "第一版运行范围停留在 CAM 项目与 dry-run，不生成机床可用 NC。",
    )

    add_heading(doc, "插件中心与诊断统计", level=2)
    add_table(
        doc,
        ["模块", "用户看到", "首版操作"],
        [
            ["插件中心", "已安装/可安装、用途、依赖、适配版本、更新时间", "安装、卸载、更新、查看变更、安装推荐组合"],
            ["版本适配包", "PowerMill/NX 版本矩阵、已验证样本、可用方法", "选择版本、运行连接测试、查看响应样例"],
            ["诊断页", "实例心跳、源游标、事件延迟、任务队列、错误摘要", "重连、重扫、取消任务、导出诊断包"],
            ["任务统计", "记录量、manual 比例、会话耗时、重复流程、配方复用", "按项目/版本/时间比较，导出报表"],
        ],
        [2100, 4350, 2910],
    )

    add_heading(doc, "黄岩工厂现场试点", level=1, page_break_before=True)
    add_body(
        doc,
        "现场验证先覆盖 5-8 家黄岩模具企业，优先选择 PowerMill/NX 版本和团队规模不同、既有数字化程度"
        "不同的样本。每家观察至少一个真实编程任务，不用演示零件替代生产复杂度。两家版本环境明确、"
        "高频流程具有代表性的企业进入 4 周产品试点。",
    )
    add_heading(doc, "参与角色与样本", level=2)
    add_table(
        doc,
        ["角色", "建议样本", "需要观察"],
        [
            ["老板/厂长", "每厂 1 人", "订单结构、交期、投资决策、软件投入"],
            ["CAM 主管", "每厂 1 人", "标准、审阅、分工、异常和版本管理"],
            ["高级程序员", "每厂 1-2 人", "复杂任务、参数理由、回滚、专家示范"],
            ["初级程序员", "每厂 1 人", "学习路径、模板依赖、常见错误和求助"],
            ["机床/质量", "每厂 1 人", "仿真、碰撞、交接、返工和问题闭环"],
            ["IT/数字化", "每厂 0-1 人", "部署、版本、性能、网络、现有系统接口"],
        ],
        [2300, 1800, 5260],
    )

    add_heading(doc, "观察与测量协议", level=2)
    protocol = [
        "先做 2 周基线：不改变工作方式，只记录任务类型、开始/结束、手工动作、自动化动作、等待和回滚。",
        "对每个任务记录 PowerMill/NX 版本、项目规模、刀路数量、任务复杂度和操作者资历，避免简单件与复杂件直接比较。",
        "采集宏、日志、Journal 和结构化事件，并记录它们与项目、实例和任务的映射质量。",
        "由高级程序员口述关键参数和异常处理理由，标记哪些动作可以标准化、哪些必须保留判断。",
        "4 周试点使用匹配任务或同一模具家族前后比较；系统自动记录使用频率、节省步骤、响应和回滚。",
        "每周复盘误识别、漏事件、错误会话归属和配方不可用原因；错误样本进入回归测试。",
    ]
    for item in protocol:
        add_numbered(doc, item)

    add_heading(doc, "KPI 与假设目标", level=2)
    add_callout(
        doc,
        "不得提前承诺",
        "下表的改善幅度是产品假设，不是行业事实或销售承诺。只有在同类任务、同一版本、明确基线和"
        "样本量说明下才能对外引用。",
        fill=CAUTION_FILL,
        color=CAUTION_TEXT,
    )
    add_table(
        doc,
        ["指标", "定义", "首轮假设", "通过条件"],
        [
            ["CAM 编程用时", "从项目准备到可提交仿真的有效时间", "降低 20%-35%", "两家样板厂均有同向改善"],
            ["重复手工动作", "已识别高频流程中的 manual 动作数", "降低 30%-50%", "不增加错误/回滚率"],
            ["配方复用率", "满足条件任务中使用已审阅配方的比例", "达到 30%+", "至少 3 类流程持续复用"],
            ["回滚/重做率", "撤销、重算、重新导入或主管退回次数", "降低 10%-20%", "按任务复杂度校正"],
            ["交接时间", "程序员提交到机床/质量完成审阅", "降低 20%", "任务状态完整率 100%"],
            ["实例映射准确率", "事件与 PowerMill/NX 窗口、项目的匹配比例", "达到 99%+", "未映射可在一处校正"],
            ["性能", "采集开销、查询延迟、宿主超时", "符合工程预算", "目标项目上 p95 达标"],
        ],
        [1900, 3350, 1900, 2210],
    )

    add_heading(doc, "Go / No-Go 评审", level=2)
    go_no_go = [
        "Go：三类以上流程在两家样板厂复用；用户持续使用自动记录；连接、会话归属和 dry-run 结果可靠。",
        "Conditional Go：效率有改善但版本适配成本高，缩小到一个产品/两个版本继续。",
        "No-Go：结构化数据覆盖不足、配方复用率低、实例映射长期不准、或收益主要来自一次性顾问配置。",
    ]
    for item in go_no_go:
        add_bullet(doc, item)

    add_heading(doc, "路线图、团队与商业假设", level=1, page_break_before=True)
    add_heading(doc, "分阶段路线", level=2)
    add_table(
        doc,
        ["阶段", "时间", "产品结果", "退出条件"],
        [
            ["0 离线底座", "已完成", "日志/Journal 学习、插件内核、dry-run、Codex 桥", "当前仓库测试通过"],
            ["1 只读连接", "0-3 个月", "PowerMill/NX 版本化查询、自动结构化采集", "多实例不串线、性能达标"],
            ["2 目标 dry-run", "3-6 个月", "测试项目中按实例运行已审阅命令和配方", "响应、差异、仿真和回滚验收"],
            ["3 工厂工艺插件", "6-12 个月", "钻孔、粗加工、精加工、批改参数等高频模块", "稳定复用、工厂付费续约"],
            ["4 模板生态", "12 个月后", "场景模板、版本适配包和区域服务商市场", "模板复用、维护成本、伙伴收入"],
        ],
        [1600, 1450, 3900, 2410],
    )
    add_heading(doc, "建议团队", level=2)
    add_table(
        doc,
        ["角色", "投入", "责任"],
        [
            ["产品/模具工艺负责人", "1.0", "现场研究、优先级、工艺审阅、试点收益"],
            ["Python/平台工程师", "1.0", "事件、采集、存储、学习、服务与性能"],
            ["PowerMill .NET 工程师", "1.0", "COM/API、PluginFramework、版本适配"],
            ["NXOpen 工程师", "1.0", "Journal、Python stubs、稳定对象查询、目标版本"],
            ["QA/应用工程师", "1.0", "项目样本、仿真、回归和现场验收"],
            ["UX/交付工程师", "各 0.2-0.4", "单窗口交互、插件安装、部署、培训和诊断"],
        ],
        [2800, 1200, 5360],
    )
    add_heading(doc, "成本与收费假设", level=2)
    add_callout(
        doc,
        "假设口径",
        "以下为内部规划区间，不是供应商报价或市场验证结果；不含 PowerMill/NX 许可证、机床停机、差旅、"
        "专用服务器和第三方仿真费用。",
        fill=CAUTION_FILL,
        color=CAUTION_TEXT,
    )
    add_table(
        doc,
        ["项目", "规划区间", "说明"],
        [
            ["12 周首版与两家样板试点", "人民币 45-75 万元", "5 人核心团队 + 兼职 UX/交付；以现有底座继续开发"],
            ["3-6 个月 live 适配与更多版本", "人民币 80-150 万元", "取决于版本数量、许可环境和现场验证成本"],
            ["工厂试点收费假设", "人民币 8-15 万元/厂/8 周", "包含部署、基线、培训和有限适配；需访谈验证"],
            ["年订阅假设", "人民币 12-30 万元/站点", "基础平台 + 指定适配包；定制工艺插件另计"],
        ],
        [3200, 2400, 3760],
    )
    add_heading(doc, "商业切入", level=2)
    add_body(
        doc,
        "优先客户画像不是“所有模具企业”，而是拥有 3 个以上 CAM 席位、PowerMill/NX 中至少一种为核心、"
        "存在高级/初级程序员梯队、重复模具家族明显、愿意开放日志样本和现场观察的企业。第一单销售的"
        "核心不是 AI，而是“把三类重复流程变成可审阅插件，并给出前后基线”。",
    )
    add_bullet(doc, "区域编程服务商可成为渠道和模板共创方，按软件版本和模具场景共同维护插件。")
    add_bullet(doc, "总承包商/N+X 路径可采用“共性平台 N + PowerMill/NX/工艺插件 X”的交付结构。")
    add_bullet(doc, "续约取决于活跃配方、节省工时、连接稳定性和版本维护，不依赖一次性演示效果。")

    add_heading(doc, "产品实施难点", level=2)
    add_table(
        doc,
        ["难点", "影响", "产品应对"],
        [
            ["PowerMill/NX 版本差异", "接口、对象名和响应不一致", "版本适配包、目标 stubs、实机样本矩阵"],
            ["日志不等于操作意图", "重复动作可能来自返工或探索", "项目实体、会话对比、专家说明和人工标签"],
            ["后台记录性能", "大项目下影响 CAM 交互", "增量 tail、WAL、分页、退避扫描和性能仪表盘"],
            ["初次采用率", "功能多但用户不会形成日常习惯", "安装后自动工作、默认视图、推荐组合和一页主流程"],
            ["工厂定制过多", "交付周期和维护成本失控", "稳定插件接口、80% 共性模块、场景参数而非分支代码"],
            ["模板持续维护", "版本升级后配方逐渐失效", "配方版本、来源样本、回归库和批量兼容检查"],
        ],
        [3200, 1700, 4460],
    )

    add_heading(doc, "现场访谈与观察提纲", level=1, page_break_before=True)
    add_callout(
        doc,
        "使用方式",
        "每次 60-90 分钟访谈 + 90-180 分钟真实编程观察。围绕一个完整任务记录软件版本、项目规模、"
        "手工/自动步骤、等待、回滚、参数理由和可复用候选。",
    )
    add_heading(doc, "老板 / 厂长 / CAM 主管", level=2)
    questions = [
        "过去三个月最常延期的模具类型是什么？延误主要发生在编程、计算、审阅、机床等待还是返工？",
        "PowerMill 与 NX 分别有多少席位、版本和常用插件？哪些版本必须优先支持？",
        "高级程序员每周有多少时间用于重复设置、审阅初级程序、排障和带教？",
        "目前如何判断一个 CAM 程序“可以提交仿真/机床”？谁签字，证据在哪里？",
        "如果软件自动记录行为，哪些日志/API 能完整反映步骤？哪些关键动作目前没有结构化来源？",
        "愿意为什么结果付费：节省编程时间、降低返工、缩短交接、保留专家经验，还是版本维护？",
    ]
    for question in questions:
        add_numbered(doc, question)

    add_heading(doc, "高级 / 初级程序员", level=2)
    programmer_questions = [
        "请从空项目开始完成一个真实任务，并在关键参数处说明为什么这样选。",
        "哪些步骤每个项目都做？哪些只在某类材料、模具结构、刀具或机床上做？",
        "最近一次撤销、重算或返工是什么原因？日志中能看到原因吗？",
        "你使用宏、Journal、模板或个人脚本吗？哪些稳定，哪些经常因版本/对象名失效？",
        "如果系统给出重复配方，你希望先看到来源、支持度、差异、版本、参数解释中的哪三项？",
        "哪些动作适合一键运行，哪些需要逐步查看结果，哪些必须保留工程师判断？",
        "如何定义“专家示范”？你更愿意按会话、步骤还是项目家族来标记？",
    ]
    for question in programmer_questions:
        add_numbered(doc, question)

    add_heading(doc, "机床 / 质量 / IT", level=2)
    ops_questions = [
        "从 CAM 程序提交到可上机之间有哪些检查？碰撞、过切、余量、刀具和后处理分别如何流转？",
        "发生异常时如何定位到具体项目、程序员、软件版本、命令和修改？目前证据缺在哪里？",
        "工作站配置、PowerMill/NX 版本、COM 注册、Python 环境和升级窗口分别是什么？",
        "多窗口、远程桌面、Teamcenter 或共享目录会如何影响会话归属和项目识别？",
        "可接受的后台 CPU、内存、磁盘和日志保留期是多少？发生故障时希望如何暂停和恢复？",
    ]
    for question in ops_questions:
        add_numbered(doc, question)

    add_heading(doc, "观察记录表", level=2)
    add_table(
        doc,
        ["任务/步骤", "开始-结束", "来源模式", "回滚/等待", "参数理由与可自动化边界"],
        [
            ["", "", "manual / automation", "", ""],
            ["", "", "manual / automation", "", ""],
            ["", "", "manual / automation", "", ""],
            ["", "", "manual / automation", "", ""],
            ["", "", "manual / automation", "", ""],
        ],
        [2200, 1350, 1800, 1600, 2410],
    )
    add_heading(doc, "访谈结束确认", level=2)
    add_bullet(doc, "向参与者回放我们理解的流程、参数和动作含义，允许其纠正。")
    add_bullet(doc, "确认 PowerMill/NX 版本、日志/API 来源、项目类型和典型样本规模。")
    add_bullet(doc, "记录三个最高频候选流程，以及三个必须保留工程师判断的步骤。")
    add_bullet(doc, "不把单个受访者意见外推为黄岩全行业结论；保留样本角色和版本分布。")

    add_heading(doc, "验收清单", level=1, page_break_before=True)
    acceptance = [
        "研究：每个产业事实、官方接口事实和竞品公开能力都有可打开来源；访谈缺口明确标注。",
        "产品：默认零业务模块，插件可在软件内安装；记录插件安装后自动连接和记录，可分类暂停。",
        "连接：Codex/NX/PowerMill 摘要在首页可见；多实例逐项显示且不猜目标。",
        "日志：支持 manual 等来源模式筛选和 L0-L4 层级筛选；专家标签人工设置。",
        "兼容：ActivityEvent 只加可选字段；PowerMill 与 NX 解析器隔离；共享动作仍为 cam.*。",
        "命令：首版提供查询和 dry-run；任务绑定版本、目标实例、项目快照，并展示响应和差异。",
        "范围：学习日志形成配方和 CAM 项目预览，不生成机床可用 NC/postprocess。",
        "性能：后台采集、导入、查询、多实例和恢复在目标工厂项目上通过预算。",
        "试点：5-8 家调研、两家样板、2 周基线 + 4 周试点；效率数字按假设管理。",
        "发布：文档、代码、插件 manifest、Skills、测试和版本说明进入同一 GitHub 分支。",
    ]
    for item in acceptance:
        add_bullet(doc, item)

    add_heading(doc, "最终建议", level=2)
    add_body(
        doc,
        "建议启动 12 周首版，但把成功定义为“稳定捕获和复用专家流程”，而不是“AI 自动编程”。先在两家样板"
        "工厂完成 PowerMill/NX 只读连接、多窗口状态、模式/层级筛选、三类重复配方和目标实例 dry-run。"
        "只有当配方复用、连接稳定、用户持续使用且 KPI 有可重复改善后，才投资 live 执行和更多工艺插件。",
        after=0,
    )

    add_heading(doc, "来源与核验说明", level=1, page_break_before=True)
    add_body(
        doc,
        "外部页面均在 2026-08-23 直接打开核验。黄岩产业规模页面为黄岩区政府转载《经济日报》内容；"
        "竞品数字为厂商公开声明，未独立验证。OpenAI/Codex 插件交付描述以当前仓库 manifest、Skills 和"
        "本地运行结果为准，未把无法打开的搜索摘要作为证据。",
    )
    add_source(
        doc,
        "S1",
        "关于组织开展黄岩区模具行业中小企业数字化改造试点企业和总承包商申报工作的通知",
        "台州市黄岩区人民政府 / 区经信科技局",
        "2024-06-27",
        "https://www.zjhy.gov.cn/art/2024/6/27/art_1633733_59109550.html",
        "核验 2024 数字化计划、2 个以上场景、约 10 家试点、2 家以上承包商和“N+X”清单。",
    )
    add_source(
        doc,
        "S2",
        "关于黄岩区模具制造行业中小企业“N+X”数字化改造试点企业项目验收评审结果及奖补资金的公示",
        "台州市黄岩区人民政府 / 区经信科技局",
        "2024-12-02",
        "https://www.zjhy.gov.cn/art/2024/12/2/art_1621869_59114040.html",
        "核验 11 家验收、5 家样本及申请、评审、现场考察流程。",
    )
    add_source(
        doc,
        "S3",
        "浙江台州黄岩区做强模具产业集群 塑造转型升级新优势",
        "台州市黄岩区人民政府，来源《经济日报》",
        "2024-09-30",
        "https://www.zjhy.gov.cn/art/2024/9/30/art_1635827_59112120.html",
        "核验从业者、企业、产值、设计/编程服务企业和数字化覆盖等产业数据。",
    )
    add_source(
        doc,
        "S4",
        "Recording macros in PowerMill",
        "Autodesk PowerMill Help",
        "2025",
        "https://help.autodesk.com/cloudhelp/2025/ENU/PWRM-ReferenceHelp/files/GUID-8B22DBDB-509A-4E66-8AD4-D06BA6720DFC.htm",
        "核验宏录制入口及只记录用户改变的对话框值。",
    )
    add_source(
        doc,
        "S5",
        "PowerShape and PowerMill API",
        "Autodesk GitHub",
        "公开仓库",
        "https://github.com/Autodesk/PowerShapeAndPowerMillAPI",
        "核验 Autodesk 官方自动化库、宏命令任务和 MIT 许可。",
    )
    add_source(
        doc,
        "S6",
        "PowerMill API Examples",
        "Autodesk GitHub",
        "公开仓库",
        "https://github.com/Autodesk/powermill-api-examples",
        "核验 PowerMill.dll、PluginFramework.dll、COM 注册和插件启用方式；仓库注明维护优先级较低。",
    )
    add_source(
        doc,
        "S7",
        "NX Open Python Reference Guide",
        "Siemens Digital Industries Software",
        "NX 2406 Series",
        "https://docs.sw.siemens.com/en-US/doc/209349590/PL20231101866122454.custom_api.nxopen_python_ref",
        "核验官方 NXOpen Python 参考入口与版本系列。",
    )
    add_source(
        doc,
        "S8",
        "NX Open Programmer's Guide",
        "Siemens Digital Industries Software",
        "NX 2212 Series",
        "https://docs.sw.siemens.com/en-US/doc/209349590/PL20220512394070742.nxopen_prog_guide/xid1124929",
        "核验自动化创建/执行、Python、许可和 Teamcenter 等官方指南结构。",
    )
    add_source(
        doc,
        "S9",
        "CAM Assist",
        "CloudNC",
        "厂商页面",
        "https://www.cloudnc.com/zh",
        "核验与现有 CAM 集成、NX 支持、人工审阅/调整和厂商效率声明。",
    )
    add_source(
        doc,
        "S10",
        "OptiNC Plugin for PowerMill",
        "OptiNC",
        "厂商页面",
        "https://www.optinc.tech/optinc/?lang=en",
        "核验 PowerMill 插件、100+ 功能及粗精加工、钻孔、3+2、批量参数等公开能力。",
    )

    add_heading(doc, "当前仓库依据", level=2)
    repo_sources = [
        "README.md：0.5.0 插件化、本地采集、多实例、日志接口、dry-run 与运行边界。",
        "docs/architecture.md：ActivityEvent v1、PowerMill/NX 边界、PMAutomation/QueueCommand 适配形态和性能测试上限。",
        "docs/mvp.md：现有演示闭环、NX 离线流水线与下一步原生适配路线。",
        "plugins/*/app-plugin.json：插件依赖、版本、分类与可安装目录。",
        "plugins/ug-cam-copilot/skills/*：NX 与 PowerMill 的 Codex 工作流学习技能和运行规则。",
    ]
    for item in repo_sources:
        add_bullet(doc, item)

    add_callout(
        doc,
        "结论置信度",
        "产业规模、政策试点和官方接口入口为高置信度；工厂痛点排序、KPI、成本和定价为中低置信度，"
        "必须通过现场样本、版本矩阵和真实任务基线提升置信度。",
        fill=CAUTION_FILL,
        color=CAUTION_TEXT,
    )


def audit_document(doc: Document) -> None:
    section = doc.sections[0]
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.left_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.top_margin == Inches(1)
    assert section.bottom_margin == Inches(1)

    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        assert tbl_w is not None and int(tbl_w.get(qn("w:w"))) == PAGE_WIDTH_DXA
        assert tbl_ind is not None and int(tbl_ind.get(qn("w:w"))) == TABLE_INDENT_DXA
        grid_widths = [
            int(col.get(qn("w:w")))
            for col in table._tbl.tblGrid.findall(qn("w:gridCol"))
        ]
        assert sum(grid_widths) == PAGE_WIDTH_DXA
        for row in table.rows:
            assert len(row.cells) == len(grid_widths)
            for index, cell in enumerate(row.cells):
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                assert tc_w is not None
                assert int(tc_w.get(qn("w:w"))) == grid_widths[index]


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    build_content(doc)
    audit_document(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
